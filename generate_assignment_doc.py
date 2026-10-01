# -*- coding: utf-8 -*-
"""Generates Assignment No. 4 lab report docx for VIT ProjectHub, matching the
reference document's structure (cover page, theory, procedure, observations, result, conclusion)."""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Kamran Btech\Sem 5\CNT\Assigment-4-website"
LOGO = os.path.join(BASE, "public", "img", "vit-logo.png")
SHOTS = r"C:\Users\kamra\AppData\Local\Temp\claude\c--Kamran-Btech-Sem-5-CNT-Assigment-4-website\675362df-2fd0-447e-844e-09b258b63082\scratchpad\docshots"
OUT = os.path.join(BASE, "Assignment 4 Roll no. 06 TY G.docx")

doc = Document()

# ---------- Base style ----------
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)

for section in doc.sections:
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def add_page_border(section):
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'double')
        el.set(qn('w:sz'), '18')
        el.set(qn('w:space'), '24')
        el.set(qn('w:color'), '000000')
        pgBorders.append(el)
    sectPr.append(pgBorders)


add_page_border(doc.sections[0])


def para(text='', bold=False, italic=False, size=11, align=None, space_after=6, space_before=0, color=None, font='Times New Roman'):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = font
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def heading(text, size=18, space_before=16, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    return p


def subheading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    return p


def bullets(items, bold_lead=None):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'


def numbered(items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'


def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(10)


def flow_box(lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(3.2)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:color'), '1F4E79')
        borders.append(el)
    tcPr.append(borders)
    cell.paragraphs[0].text = ''
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_screenshot(filename, caption, width=Inches(6.0)):
    path = os.path.join(SHOTS, filename)
    doc.add_picture(path, width=width)
    last_p = doc.paragraphs[-1]
    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(10)
    cap.paragraph_format.space_after = Pt(14)


# ================= COVER PAGE =================
para('', size=6)
para("Bansilal Ramnath Agarwal Charitable Trust\u2019s", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run = p.add_run("Vishwakarma Institute of Technology, Pune-37")
run.bold = True
run.font.size = Pt(22)
run.font.name = 'Times New Roman'

para("(An Autonomous Institute of Savitribai Phule Pune University)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

if os.path.exists(LOGO):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_p.add_run()
    run.add_picture(LOGO, width=Inches(1.6))

para('', size=6)
para('', size=6)
para('', size=6)
para('', size=6)

info_table = doc.add_table(rows=5, cols=2)
info_table.style = 'Table Grid'
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
rows = [
    ("Division", "TY AIML - G"),
    ("Batch", "1"),
    ("PRN", "1252030006"),
    ("Roll-no", "06"),
    ("Name", "Pathan Kamran Amjad"),
]
for i, (k, v) in enumerate(rows):
    info_table.cell(i, 0).text = ''
    r = info_table.cell(i, 0).paragraphs[0].add_run(k)
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = 'Times New Roman'
    info_table.cell(i, 1).text = ''
    r2 = info_table.cell(i, 1).paragraphs[0].add_run(v)
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.name = 'Times New Roman'
    info_table.rows[i].height = Cm(1.1)

doc.add_page_break()

# ================= ASSIGNMENT CONTENT =================
para("Assignment No. 4", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

heading("Title of the Experiment", size=18)
para(
    "Design and Develop a Responsive Full-Stack Web Application \u2018VIT ProjectHub\u2019 for Managing "
    "VIT Student Projects and Assignments Using Node.js, Express.js and MongoDB Having Landing Page, "
    "Login Page, Register Page, Student Dashboard, Teacher Dashboard, Classroom Page and Assignment "
    "Submission Page",
    bold=True,
)
hr()

heading("Aim")
para(
    "To design and develop a responsive full-stack web application named VIT ProjectHub for managing "
    "student projects and assignments at Vishwakarma Institute of Technology (VIT), Pune, using HTML5, "
    "CSS3, Bootstrap 5 and jQuery for the frontend, Node.js and Express.js for the backend, and MongoDB "
    "as the database. The application should provide teacher and student registration/login, classroom "
    "creation and joining using a unique class code, assignment posting with deadlines across four "
    "academic categories \u2014 SDP, EDI, Design Thinking and Course Projects \u2014 file-based submissions, "
    "and grading with marks and feedback."
)
hr()

heading("Theory")
para(
    "VIT ProjectHub is a web-based academic platform that enables teachers to create classrooms, post "
    "assignments and grade student submissions, while students can join classrooms, submit their work and "
    "track deadlines and feedback. Modern full-stack web applications of this kind are commonly built using "
    "Node.js, Express.js and MongoDB on the backend, paired with Bootstrap and jQuery on the frontend."
)

subheading("Node.js")
para(
    "Node.js is a JavaScript runtime environment that executes JavaScript code outside the browser. It is "
    "mainly used to build scalable backend applications and APIs."
)
para("Features:", space_after=2)
bullets([
    "Non-blocking I/O",
    "Event-driven architecture",
    "High performance",
    "JavaScript on both client and server",
])

subheading("Express.js")
para("Express.js is a lightweight framework built on Node.js that simplifies backend development.")
para("Functions:", space_after=2)
bullets([
    "Create REST APIs",
    "Handle HTTP requests",
    "Routing",
    "Middleware support",
    "Database connectivity",
])

subheading("MongoDB")
para("MongoDB is a NoSQL database that stores data in JSON-like documents instead of tables.")
para("Collections used in this project:", space_after=2)
bullets(["Users", "Classrooms", "Assignments", "Submissions", "Resources", "Projects", "Notifications"])

subheading("Mongoose")
para(
    "Mongoose is an Object Data Modeling (ODM) library that connects Node.js applications with MongoDB. "
    "It provides schemas and models for structured database operations and validation."
)

subheading("Bootstrap 5")
para(
    "Bootstrap 5 is a CSS framework used to create responsive and attractive user interfaces. In this "
    "project it powers the nav-pills used for switching between the SDP, EDI, DT and Course Project "
    "categories, along with modals, toasts, tabs, badges and the responsive sidebar layout."
)

subheading("jQuery")
para(
    "jQuery is a fast, lightweight JavaScript library used to simplify DOM manipulation, event handling and "
    "AJAX-style dynamic updates. In this project jQuery drives category-pill filtering, live assignment "
    "search, client-side form validation, Bootstrap modal/toast wiring, and dynamic re-rendering of "
    "classroom tabs and dashboards without page reloads."
)

subheading("JSON Web Token (JWT)")
para(
    "JWT is a compact, URL-safe token format used to securely transmit authentication information between "
    "the client and server. On login, the server issues a signed token containing the user\u2019s id and role, "
    "which the frontend attaches to every subsequent API request to prove identity."
)

subheading("Multer")
para(
    "Multer is a Node.js middleware used for handling multipart/form-data, primarily used for uploading "
    "files. It is used here to accept assignment submissions and classroom resources, with file type and "
    "size validation."
)

subheading("REST API")
para("The frontend communicates with the backend through REST APIs.")
para("Examples:", space_after=2)
code_block(
    "POST /api/auth/login\n"
    "POST /api/classrooms\n"
    "POST /api/classrooms/join\n"
    "POST /api/assignments/classroom/:classroomId\n"
    "POST /api/submissions/assignment/:assignmentId\n"
    "PUT  /api/submissions/:id/grade"
)
hr()

heading("Objective")
bullets([
    "To understand full-stack web application development for an academic institute.",
    "To develop a responsive project and assignment management portal for VIT students and faculty.",
    "To implement role-based user authentication (student/teacher) using JWT.",
    "To perform CRUD operations using MongoDB via Mongoose.",
    "To implement REST APIs using Express.js.",
    "To use jQuery for dynamic UI updates, form validation, filtering and DOM manipulation.",
    "To implement secure file upload and validation using Multer.",
    "To manage user data and sessions securely.",
    "To create a modern, responsive interface using Bootstrap 5.",
])
hr()

heading("Apparatus / Software Required")
subheading("Hardware")
bullets(["Laptop/Desktop", "Minimum 4 GB RAM", "Intel i3 Processor or Higher", "Internet Connection"])
subheading("Software")
bullets([
    "Windows 10/11",
    "Visual Studio Code",
    "Node.js",
    "npm",
    "MongoDB Community Server",
    "MongoDB Compass",
    "Express.js",
    "Bootstrap 5",
    "jQuery",
    "JSON Web Token (jsonwebtoken)",
    "Multer",
    "Google Chrome",
    "Git (Optional)",
])
hr()

heading("Modules Developed")
para("VIT ProjectHub contains the following modules:")
numbered([
    "Landing Page",
    "Registration Page",
    "Login Page",
    "Student Dashboard",
    "Teacher Dashboard",
    "Classroom Page (Stream, Assignments, Resources, Projects, People tabs)",
    "Assignment Details Page",
    "Submission / Review (Grading) Page",
])
hr()

heading("Database Collections")

subheading("User Collection")
para("Stores:", space_after=2)
bullets(["Name", "Email", "Password (Encrypted)", "Role (Student / Teacher)", "Registration Number", "Department"])

subheading("Classroom Collection")
para("Stores:", space_after=2)
bullets(["Name", "Subject", "Category (SDP / EDI / DT / Course Project)", "Unique Class Code", "Teacher", "Enrolled Students"])

subheading("Assignment Collection")
para("Stores:", space_after=2)
bullets(["Classroom", "Title", "Description", "Category", "Stage", "Max Marks", "Deadline", "Attached Resources"])

subheading("Submission Collection")
para("Stores:", space_after=2)
bullets(["Assignment", "Student", "Uploaded Files", "Note", "Status", "Submitted Date", "Marks", "Feedback", "Graded By"])

subheading("Resource Collection")
para("Stores:", space_after=2)
bullets(["Classroom", "Assignment (optional)", "Title", "Description", "File", "Uploaded By"])

subheading("Project Collection")
para("Stores:", space_after=2)
bullets(["Classroom", "Title", "Problem Statement", "Category", "Current Stage", "Team Members", "Mentor", "Status"])

subheading("Notification Collection")
para("Stores:", space_after=2)
bullets(["User", "Message", "Type", "Link", "Read Status"])
hr()

heading("Project Architecture")
flow_box([
    "Browser", "\u2193",
    "HTML / CSS / Bootstrap 5 / jQuery (Frontend)", "\u2193",
    "Fetch API (AJAX Requests)", "\u2193",
    "Express Server", "\u2193",
    "Middleware (Auth, Role, Multer)", "\u2193",
    "Controllers", "\u2193",
    "Mongoose Models", "\u2193",
    "MongoDB Database",
])
hr()

heading("Procedure / Algorithm")
steps = [
    ("Step 1", "Install Node.js.", None),
    ("Step 2", "Create Project Folder and Initialize Node Project.", "npm init -y"),
    ("Step 3", "Install Backend Packages.", "npm install express mongoose cors dotenv bcryptjs jsonwebtoken multer express-validator"),
    ("Step 4", "Install Development Dependency.", "npm install --save-dev nodemon"),
    ("Step 5", "Configure Environment Variables.\nCreate:", ".env\n\nStore:\nMONGO_URI=your_database_connection_string\nJWT_SECRET=your_secret_key\nPORT=5000"),
    ("Step 6", "Create Mongoose Models.", None),
    ("Step 7", "Create Middleware.", None),
    ("Step 8", "Create Controllers.", None),
    ("Step 9", "Create Routes.", None),
    ("Step 10", "Configure server.js to connect MongoDB, mount API routes and serve static frontend files.", None),
    ("Step 11", "Develop Frontend Pages.", None),
    ("Step 12", "Add Bootstrap 5 nav-pills for SDP / EDI / DT / Course Project category switching and jQuery for dynamic filtering, modals, toasts and form validation.", None),
    ("Step 13", "Connect Frontend with Backend using the Fetch API.", None),
    ("Step 14", "Run MongoDB Server (local installation or Docker container).", None),
    ("Step 15", "Seed Demo Data.", "npm run seed"),
    ("Step 16", "Run Backend Server.", "npm run dev"),
    ("Step 17", "Open Browser.", "http://localhost:5000"),
]
model_lists = {
    "Step 6": ["User Model", "Classroom Model", "Assignment Model", "Submission Model", "Resource Model", "Project Model", "Notification Model"],
    "Step 7": ["auth.js \u2013 JWT verification", "role.js \u2013 role-based access guard", "upload.js \u2013 Multer file upload + validation", "errorHandler.js \u2013 centralized error handling"],
    "Step 8": ["Authentication Controller", "Classroom Controller", "Assignment Controller", "Submission Controller", "Resource Controller", "Project Controller", "Notification Controller"],
    "Step 9": ["Authentication Routes", "Classroom Routes", "Assignment Routes", "Submission Routes", "Resource Routes", "Project Routes", "Notification Routes"],
    "Step 11": ["Landing Page", "Login Page", "Register Page", "Student Dashboard", "Teacher Dashboard", "Classroom Page", "Assignment Page", "Submission Page"],
}

for label, text, code in steps:
    subheading(label, size=12)
    para(text)
    if code:
        code_block(code)
    if label in model_lists:
        bullets(model_lists[label])
hr()

heading("Flowchart")
flow_box([
    "Start", "\u2193",
    "Open Website", "\u2193",
    "Landing Page", "\u2193",
    "Register / Login", "\u2193",
    "Role Check", "\u2193",
    "Teacher Dashboard  /  Student Dashboard", "\u2193",
    "Create Classroom  /  Join Classroom with Code", "\u2193",
    "Post Assignment  /  View Assignment", "\u2193",
    "Upload Submission", "\u2193",
    "Review Submission & Grade", "\u2193",
    "View Grade & Feedback", "\u2193",
    "End",
])
hr()

heading("Features Implemented")
bullets([
    "Responsive Design",
    "Teacher and Student Registration",
    "Role-based Login",
    "JWT-based Secure Authentication",
    "Classroom Creation with Unique Class Code",
    "Join Classroom Using Class Code",
    "Dynamic SDP / EDI / DT / Course Project Category Pills",
    "Assignment Posting with Deadlines and Stages",
    "Resource Upload with File Validation",
    "Multi-file Assignment Submission",
    "Assignment Search and Filter",
    "Submission Status Tracking (Pending / Submitted / Late / Missed / Graded)",
    "Grading with Marks and Feedback",
    "Real-time Notifications",
    "Bootstrap Modals, Toasts and Tabs",
    "Responsive Sidebar Navigation",
])
hr()

heading("Observations")
obs_table = doc.add_table(rows=1, cols=3)
obs_table.style = 'Table Grid'
hdr = obs_table.rows[0].cells
hdr[0].text = 'Sr. No.'
hdr[1].text = 'Observation'
hdr[2].text = 'Status'
for c in hdr:
    for r in c.paragraphs[0].runs:
        r.bold = True

observations = [
    "Landing page displayed successfully with SDP / EDI / DT / Course Project pills",
    "User registration completed successfully for both roles",
    "User login authenticated successfully using JWT",
    "Teacher able to create classroom and generate unique class code",
    "Student able to join classroom using class code",
    "Assignment created successfully by teacher with deadline and stage",
    "Resource uploaded successfully with file type validation",
    "Student able to submit assignment with multiple files",
    "Submission status updates correctly (Pending / Submitted / Late / Missed)",
    "Teacher able to view submission roster and grade with marks and feedback",
    "Student able to view grade and feedback after grading",
    "Notifications generated for classroom join, new assignment and grading",
    "Category pills dynamically filter classrooms and assignments",
    "Responsive layout verified on desktop and mobile screen sizes",
    "MongoDB stores all user, classroom, assignment and submission data successfully",
]
for i, obs in enumerate(observations, start=1):
    row = obs_table.add_row().cells
    row[0].text = str(i)
    row[1].text = obs
    row[2].text = '\u2714'
doc.add_paragraph().paragraph_format.space_after = Pt(10)
hr()

heading("Advantages")
bullets([
    "Centralized platform for all academic project categories.",
    "Role-based secure access for students and faculty.",
    "Real-time tracking of deadlines and submission status.",
    "Easy classroom management using unique class codes.",
    "Structured grading and feedback workflow.",
    "Responsive user interface across devices.",
    "Scalable REST API architecture.",
    "Reusable middleware and controller structure.",
    "Efficient database operations using Mongoose.",
    "Reduces manual, paper-based assignment tracking.",
])
hr()

heading("Applications")
bullets([
    "College / Institute Project Management Portals",
    "Academic Assignment Submission Systems",
    "Learning Management Systems (LMS)",
    "Faculty Grading and Feedback Systems",
    "Classroom Management Tools",
    "Student Progress Tracking Systems",
    "Capstone / Final Year Project Tracking",
    "Online Course Assignment Platforms",
])
hr()

heading("Result")
add_screenshot("01-landing.png", "Fig. 1: Landing page")
add_screenshot("02-categories.png", "Fig. 2: SDP / EDI / DT / Course Project category pills")
add_screenshot("04-login.png", "Fig. 3: Login page")
add_screenshot("03-register.png", "Fig. 4: Registration page")
add_screenshot("05-teacher-dashboard.png", "Fig. 5: Teacher dashboard")
add_screenshot("06-classroom-stream.png", "Fig. 6: Classroom \u2013 Stream tab")
add_screenshot("07-classroom-assignments.png", "Fig. 7: Classroom \u2013 Assignments tab")
add_screenshot("08-classroom-people.png", "Fig. 8: Classroom \u2013 People tab")
add_screenshot("09-assignment-roster.png", "Fig. 9: Assignment page \u2013 submission roster (teacher view)")
add_screenshot("10-submission-grading.png", "Fig. 10: Grading a submission with marks and feedback")
add_screenshot("11-student-dashboard.png", "Fig. 11: Student dashboard")
add_screenshot("12-student-assignment.png", "Fig. 12: Student view of a graded assignment")
hr()

heading("Conclusion")
para(
    "The VIT ProjectHub web application was successfully designed and developed using HTML5, CSS3, "
    "Bootstrap 5, jQuery, Node.js, Express.js and MongoDB. The application provides teacher and student "
    "registration, role-based login, classroom creation and joining through a unique class code, assignment "
    "posting across the SDP, EDI, Design Thinking and Course Project categories, file-based submissions, "
    "and grading with marks and feedback. The website is responsive, interactive, and demonstrates the "
    "implementation of a complete full-stack web application for academic project and assignment management."
)

doc.save(OUT)
print("Saved:", OUT)
